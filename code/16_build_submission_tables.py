#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import csv
import math
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import (
    WD_TABLE_ALIGNMENT,
    WD_CELL_VERTICAL_ALIGNMENT,
)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(".")
RESULTS = ROOT / "results"
TABLES = RESULTS / "tables"
SUBMISSION = RESULTS / "submission"

SUBMISSION.mkdir(parents=True, exist_ok=True)


# ============================================================
# General utilities
# ============================================================

def read_csv(path):
    path = Path(path)

    if not path.exists():
        raise FileNotFoundError(f"Missing input file: {path}")

    with path.open(
        "r",
        encoding="utf-8-sig",
        newline=""
    ) as f:
        return list(csv.DictReader(f))


def as_float(value):
    try:
        return float(value)
    except (TypeError, ValueError):
        return float("nan")


def fmt_n(value):
    value = as_float(value)

    if math.isnan(value):
        return "—"

    return f"{int(round(value)):,}"


def fmt_num(value, digits=2):
    value = as_float(value)

    if math.isnan(value):
        return "—"

    return f"{value:.{digits}f}"


def fmt_p(value):
    value = as_float(value)

    if math.isnan(value):
        return "—"

    if value < 0.0001:
        return "<0.0001"

    if value < 0.001:
        return f"{value:.4f}"

    return f"{value:.3f}"


def fmt_hr(row, prefix=""):
    hr = as_float(row[f"{prefix}hazard_ratio"])
    lo = as_float(row[f"{prefix}lower_95"])
    hi = as_float(row[f"{prefix}upper_95"])

    return f"{hr:.2f} ({lo:.2f}–{hi:.2f})"


def fmt_hr_sd(row):
    hr = as_float(row["hazard_ratio_per_sd"])
    lo = as_float(row["lower_95_per_sd"])
    hi = as_float(row["upper_95_per_sd"])

    return f"{hr:.2f} ({lo:.2f}–{hi:.2f})"


TRANSITION_LABELS = {
    "0_to_1_stroke": "Stroke-free to stroke",
    "0_to_2_death": "Stroke-free to death",
    "1_to_2_death": "Stroke to death",
}

MODEL_LABELS = {
    "Model_1": "Model 1",
    "Model_2": "Model 2",
    "Model_3": "Model 3",
}

SENSITIVITY_LABELS = {
    "Primary_complete_9_domains":
        "Primary: complete nine-domain score",
    "Survey_weighted_complete_9_domains":
        "Survey-weighted complete-case analysis",
    "Available_score_7plus_domains":
        "Available-domain score (≥7 domains)",
    "Complete_8_domains_without_activity":
        "Complete eight-domain score excluding activity",
}

CATEGORY_LABELS = {
    "Q1_0_to_10": "0–10 (reference)",
    "Q2_11_to_12": "11–12",
    "Q3_13_to_14": "13–14",
    "Q4_15_to_18": "15–18",
}


# ============================================================
# Word formatting
# ============================================================

def configure_document(document, title):
    section = document.sections[0]

    section.orientation = WD_ORIENT.LANDSCAPE

    old_width = section.page_width
    old_height = section.page_height

    section.page_width = old_height
    section.page_height = old_width

    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(8.5)

    normal._element.rPr.rFonts.set(
        qn("w:eastAsia"),
        "Arial"
    )

    document.core_properties.title = title
    document.core_properties.subject = (
        "Adapted Brain Care Score, stroke, and mortality"
    )
    document.core_properties.author = "Study investigators"


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()

    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")

    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()

    shading = tc_pr.find(qn("w:shd"))

    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)

    shading.set(qn("w:fill"), fill)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()

    borders = tc_pr.first_child_found_in(
        "w:tcBorders"
    )

    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)

    for edge in (
        "top",
        "left",
        "bottom",
        "right",
        "insideH",
        "insideV",
    ):
        if edge not in kwargs:
            continue

        edge_data = kwargs.get(edge)

        tag = f"w:{edge}"
        element = borders.find(qn(tag))

        if element is None:
            element = OxmlElement(tag)
            borders.append(element)

        for key in (
            "val",
            "sz",
            "space",
            "color",
        ):
            if key in edge_data:
                element.set(
                    qn(f"w:{key}"),
                    str(edge_data[key])
                )


def clear_cell_borders(cell):
    border = {
        "val": "nil",
        "sz": "0",
        "color": "FFFFFF",
    }

    set_cell_border(
        cell,
        top=border,
        bottom=border,
        left=border,
        right=border,
    )


def set_cell_margins(
    cell,
    top=50,
    start=65,
    bottom=50,
    end=65,
):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()

    tc_mar = tc_pr.first_child_found_in(
        "w:tcMar"
    )

    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)

    for margin_name, margin_value in {
        "top": top,
        "start": start,
        "bottom": bottom,
        "end": end,
    }.items():
        node = tc_mar.find(
            qn(f"w:{margin_name}")
        )

        if node is None:
            node = OxmlElement(
                f"w:{margin_name}"
            )
            tc_mar.append(node)

        node.set(qn("w:w"), str(margin_value))
        node.set(qn("w:type"), "dxa")


def format_cell(
    cell,
    font_size=8.2,
    bold=False,
    alignment=WD_ALIGN_PARAGRAPH.CENTER,
):
    cell.vertical_alignment = (
        WD_CELL_VERTICAL_ALIGNMENT.CENTER
    )

    set_cell_margins(cell)

    for paragraph in cell.paragraphs:
        paragraph.alignment = alignment
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0

        for run in paragraph.runs:
            run.font.name = "Arial"
            run.font.size = Pt(font_size)
            run.font.bold = bold

            run._element.rPr.rFonts.set(
                qn("w:eastAsia"),
                "Arial"
            )


def add_title(document, text):
    paragraph = document.add_paragraph()

    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(5)

    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(10.5)

    return paragraph


def add_panel_title(document, text):
    paragraph = document.add_paragraph()

    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(3)

    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(9)

    return paragraph


def add_note(document, text):
    paragraph = document.add_paragraph()

    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0

    label = paragraph.add_run("Note: ")
    label.bold = True
    label.font.name = "Arial"
    label.font.size = Pt(7.8)

    note = paragraph.add_run(text)
    note.font.name = "Arial"
    note.font.size = Pt(7.8)

    return paragraph


def add_three_line_table(
    document,
    headers,
    rows,
    first_column_left=True,
):
    table = document.add_table(
        rows=1,
        cols=len(headers)
    )

    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    prevent_row_split(header_row)

    top_border = {
        "val": "single",
        "sz": "12",
        "space": "0",
        "color": "000000",
    }

    middle_border = {
        "val": "single",
        "sz": "8",
        "space": "0",
        "color": "000000",
    }

    bottom_border = {
        "val": "single",
        "sz": "12",
        "space": "0",
        "color": "000000",
    }

    for index, header in enumerate(headers):
        cell = header_row.cells[index]
        cell.text = str(header)

        clear_cell_borders(cell)

        set_cell_border(
            cell,
            top=top_border,
            bottom=middle_border,
        )

        set_cell_shading(cell, "F2F2F2")

        format_cell(
            cell,
            font_size=8.3,
            bold=True,
            alignment=(
                WD_ALIGN_PARAGRAPH.LEFT
                if index == 0 and first_column_left
                else WD_ALIGN_PARAGRAPH.CENTER
            )
        )

    for row_data in rows:
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])

        for index, value in enumerate(row_data):
            cells[index].text = (
                "" if value is None else str(value)
            )

            clear_cell_borders(cells[index])

            format_cell(
                cells[index],
                font_size=8.1,
                bold=False,
                alignment=(
                    WD_ALIGN_PARAGRAPH.LEFT
                    if index == 0 and first_column_left
                    else WD_ALIGN_PARAGRAPH.CENTER
                )
            )

    if len(table.rows) > 1:
        for cell in table.rows[-1].cells:
            set_cell_border(
                cell,
                bottom=bottom_border,
            )

    return table


def add_page_break(document):
    document.add_page_break()


# ============================================================
# Main tables
# ============================================================

main_document = Document()

configure_document(
    main_document,
    "Main Tables"
)

hrs_baseline = read_csv(
    TABLES / "Table_1A_HRS_baseline.csv"
)

charls_baseline = read_csv(
    TABLES / "Table_1B_CHARLS_baseline.csv"
)

baseline_headers = [
    "Characteristic",
    "Overall",
    "0–10",
    "11–12",
    "13–14",
    "15–18",
]

add_title(
    main_document,
    "Table 1. Baseline characteristics according to adapted hBCS category"
)

add_panel_title(
    main_document,
    "Panel A. Health and Retirement Study"
)

hrs_rows = [
    [
        row["characteristic"],
        row["Overall"],
        row["Q1_0_to_10"],
        row["Q2_11_to_12"],
        row["Q3_13_to_14"],
        row["Q4_15_to_18"],
    ]
    for row in hrs_baseline
]

add_three_line_table(
    main_document,
    baseline_headers,
    hrs_rows,
)

add_panel_title(
    main_document,
    "Panel B. China Health and Retirement Longitudinal Study"
)

charls_rows = [
    [
        row["characteristic"],
        row["Overall"],
        row["Q1_0_to_10"],
        row["Q2_11_to_12"],
        row["Q3_13_to_14"],
        row["Q4_15_to_18"],
    ]
    for row in charls_baseline
]

add_three_line_table(
    main_document,
    baseline_headers,
    charls_rows,
)

add_note(
    main_document,
    "Values are mean (SD) or n (%). Baseline was defined as each "
    "participant's first eligible observation with a complete nine-domain "
    "adapted hBCS. Category cut points were derived in HRS and applied "
    "unchanged in CHARLS. BMI values outside 12–60 kg/m² were treated as "
    "missing. The CHARLS education variable is reported as a harmonized "
    "category code and should not be interpreted as years of schooling."
)

add_page_break(main_document)

primary_models = read_csv(
    RESULTS / "primary_transition_model_summary.csv"
)

cohort_rank = {
    "HRS": 1,
    "CHARLS": 2,
}

transition_rank = {
    "0_to_1_stroke": 1,
    "0_to_2_death": 2,
    "1_to_2_death": 3,
}

model_rank = {
    "Model_1": 1,
    "Model_2": 2,
    "Model_3": 3,
}

primary_models.sort(
    key=lambda row: (
        cohort_rank[row["cohort"]],
        transition_rank[row["transition"]],
        model_rank[row["model"]],
    )
)

table2_rows = []
previous_cohort = None
previous_transition = None

for row in primary_models:
    cohort = row["cohort"]
    transition = row["transition"]

    cohort_display = (
        cohort
        if cohort != previous_cohort
        else ""
    )

    transition_display = (
        TRANSITION_LABELS[transition]
        if (
            cohort != previous_cohort
            or transition != previous_transition
        )
        else ""
    )

    table2_rows.append([
        cohort_display,
        transition_display,
        MODEL_LABELS[row["model"]],
        fmt_n(row["model_intervals"]),
        fmt_n(row["model_events"]),
        fmt_hr(row),
        fmt_p(row["p_value"]),
        fmt_hr_sd(row),
        fmt_num(row["concordance"], 3),
    ])

    previous_cohort = cohort
    previous_transition = transition

add_title(
    main_document,
    "Table 2. Associations between adapted hBCS and transition-specific hazards"
)

add_three_line_table(
    main_document,
    [
        "Cohort",
        "Transition",
        "Model",
        "Intervals",
        "Events",
        "HR per 3-point increase (95% CI)",
        "P value",
        "HR per SD increase (95% CI)",
        "C-index",
    ],
    table2_rows,
)

add_note(
    main_document,
    "HR indicates hazard ratio; CI, confidence interval; SD, standard "
    "deviation. Model 1 included adapted hBCS only. Model 2 additionally "
    "adjusted for age and sex. Model 3 additionally adjusted for education, "
    "race, and ethnicity in HRS and for education category in CHARLS. "
    "All models were stratified by survey wave and used participant-clustered "
    "robust standard errors. HRs below 1 indicate that a higher adapted hBCS "
    "was associated with a lower transition hazard. CHARLS stroke-to-death "
    "models were not estimated because of the limited number of events."
)

main_path = SUBMISSION / "Main_Tables.docx"
main_document.save(main_path)


# ============================================================
# Supplementary tables
# ============================================================

supplement_document = Document()

configure_document(
    supplement_document,
    "Supplementary Tables"
)

sensitivity = read_csv(
    RESULTS / "sensitivity_model_summary.csv"
)

sensitivity.sort(
    key=lambda row: (
        cohort_rank[row["cohort"]],
        transition_rank[row["transition"]],
        list(SENSITIVITY_LABELS).index(
            row["sensitivity_analysis"]
        ),
    )
)

s1_rows = []
previous_cohort = None
previous_transition = None

for row in sensitivity:
    cohort = row["cohort"]
    transition = row["transition"]

    s1_rows.append([
        (
            cohort
            if cohort != previous_cohort
            else ""
        ),
        (
            TRANSITION_LABELS[transition]
            if (
                cohort != previous_cohort
                or transition != previous_transition
            )
            else ""
        ),
        SENSITIVITY_LABELS[
            row["sensitivity_analysis"]
        ],
        fmt_n(row["model_intervals"]),
        fmt_n(row["model_events"]),
        (
            f'{as_float(row["hazard_ratio_per_3_points"]):.2f} '
            f'({as_float(row["lower_95"]):.2f}–'
            f'{as_float(row["upper_95"]):.2f})'
        ),
        (
            f'{as_float(row["hazard_ratio_per_sd"]):.2f} '
            f'({as_float(row["lower_95_per_sd"]):.2f}–'
            f'{as_float(row["upper_95_per_sd"]):.2f})'
        ),
        row.get("p_label", "<0.0001"),
    ])

    previous_cohort = cohort
    previous_transition = transition

add_title(
    supplement_document,
    "Table S1. Sensitivity analyses of adapted hBCS and transition-specific hazards"
)

add_three_line_table(
    supplement_document,
    [
        "Cohort",
        "Transition",
        "Sensitivity analysis",
        "Intervals",
        "Events",
        "HR per 3-point increase (95% CI)",
        "HR per SD increase (95% CI)",
        "P value",
    ],
    s1_rows,
)

add_note(
    supplement_document,
    "All sensitivity models used the fully adjusted covariate set. "
    "The available-domain score required at least seven observed domains "
    "and was rescaled to 0–18. The eight-domain analysis excluded physical "
    "activity because activity was administered to a subsample in early "
    "CHARLS waves."
)

add_page_break(supplement_document)

category_models = read_csv(
    RESULTS / "hbcs_category_model_summary.csv"
)

category_rank = {
    "Q1_0_to_10": 1,
    "Q2_11_to_12": 2,
    "Q3_13_to_14": 3,
    "Q4_15_to_18": 4,
}

category_models.sort(
    key=lambda row: (
        cohort_rank[row["cohort"]],
        transition_rank[row["transition"]],
        category_rank[row["category"]],
    )
)

s2_rows = []
previous_cohort = None
previous_transition = None

for row in category_models:
    cohort = row["cohort"]
    transition = row["transition"]
    first_transition_row = (
        cohort != previous_cohort
        or transition != previous_transition
    )

    if row["category"] == "Q1_0_to_10":
        hr_text = "1.00 (reference)"
        p_value = "—"
    else:
        hr_text = fmt_hr(row)
        p_value = fmt_p(row["p_value"])

    s2_rows.append([
        cohort if cohort != previous_cohort else "",
        (
            TRANSITION_LABELS[transition]
            if first_transition_row
            else ""
        ),
        CATEGORY_LABELS[row["category"]],
        fmt_n(row["model_intervals"]),
        fmt_n(row["model_events"]),
        hr_text,
        p_value,
        (
            row["p_trend_label"]
            if first_transition_row
            else ""
        ),
    ])

    previous_cohort = cohort
    previous_transition = transition

add_title(
    supplement_document,
    "Table S2. Associations across adapted hBCS categories"
)

add_three_line_table(
    supplement_document,
    [
        "Cohort",
        "Transition",
        "Adapted hBCS category",
        "Intervals",
        "Events",
        "HR (95% CI)",
        "P value",
        "P for trend",
    ],
    s2_rows,
)

add_note(
    supplement_document,
    "The 0–10 category was the reference. Category cut points were based "
    "on the HRS distribution and applied unchanged in CHARLS. Models used "
    "the fully adjusted covariate set, survey-wave strata, and "
    "participant-clustered robust standard errors."
)

add_page_break(supplement_document)

rcs_models = read_csv(
    RESULTS / "rcs_model_summary.csv"
)

rcs_models.sort(
    key=lambda row: (
        cohort_rank[row["cohort"]],
        transition_rank[row["transition"]],
    )
)

s3_rows = []

for row in rcs_models:
    s3_rows.append([
        row["cohort"],
        TRANSITION_LABELS[row["transition"]],
        fmt_n(row["model_intervals"]),
        fmt_n(row["model_events"]),
        fmt_num(row["reference_score"], 0),
        ", ".join([
            fmt_num(row["knot_1"], 0),
            fmt_num(row["knot_2"], 0),
            fmt_num(row["knot_3"], 0),
            fmt_num(row["knot_4"], 0),
        ]),
        fmt_p(row["p_overall"]),
        fmt_p(row["p_nonlinear"]),
    ])

add_title(
    supplement_document,
    "Table S3. Restricted cubic spline model specifications and tests"
)

add_three_line_table(
    supplement_document,
    [
        "Cohort",
        "Transition",
        "Intervals",
        "Events",
        "Reference score",
        "Knot locations",
        "P overall",
        "P nonlinear",
    ],
    s3_rows,
)

add_note(
    supplement_document,
    "Restricted cubic spline models used four knots located at the 5th, "
    "35th, 65th, and 95th percentiles of the transition-specific adapted "
    "hBCS distribution. Hazard ratios were referenced to an adapted hBCS "
    "of 12. P values were obtained from robust Wald tests using the same "
    "covariate adjustment, survey-wave strata, and participant clustering "
    "as the primary models."
)

add_page_break(supplement_document)

scoring_rows = [
    [
        "Smoking",
        "0: current; 1: former; 2: never",
        "0: current; 1: former; 2: never",
        "Harmonized current and lifetime smoking status",
    ],
    [
        "Alcohol",
        "Sex-specific estimated weekly intake: 2 for ≤14 drinks/week in men or ≤7 in women; 1 for up to twice these limits; 0 above twice the limits",
        "Frequency and amount categories converted to approximate weekly intake; same sex-specific thresholds",
        "Approximate intake used because questionnaire formats differed",
    ],
    [
        "Physical activity",
        "2: frequent vigorous/moderate activity; 1: less frequent vigorous/moderate or regular light activity; 0: inactive",
        "2: vigorous/moderate activity ≥2 days/week; 1: 1 day/week or light activity ≥3 days/week; 0: otherwise",
        "Adapted to harmonized activity frequency measures",
    ],
    [
        "BMI",
        "0: <18.5 or ≥30; 1: 25–<30; 2: 18.5–<25 kg/m²",
        "Same thresholds",
        "BMI outside 12–60 kg/m² treated as missing",
    ],
    [
        "Hypertension",
        "0: ever diagnosed; 2: never diagnosed",
        "Same definition",
        "No intermediate score",
    ],
    [
        "Diabetes",
        "0: ever diagnosed; 2: never diagnosed",
        "Same definition",
        "No intermediate score",
    ],
    [
        "Sleep",
        "0: restless sleep present; 2: absent",
        "0: restless sleep ≥3 days/week; 2: <3 days/week",
        "No intermediate score",
    ],
    [
        "Mood",
        "Six CES-D items excluding sleep and loneliness; 2: symptom proportion 0; 1: >0–<0.5; 0: ≥0.5; ≥5 items required",
        "Eight CES-D items excluding sleep and loneliness; identical proportional thresholds; ≥7 items required",
        "Sleep and loneliness excluded to prevent double counting",
    ],
    [
        "Social connection",
        "0: not married/partnered; 2: married/partnered",
        "Same definition",
        "Marital/partnership status used for cross-cohort comparability",
    ],
]

add_title(
    supplement_document,
    "Table S4. Construction of the adapted harmonized Brain Care Score"
)

add_three_line_table(
    supplement_document,
    [
        "Domain",
        "HRS scoring",
        "CHARLS scoring",
        "Harmonization note",
    ],
    scoring_rows,
)

add_note(
    supplement_document,
    "Each domain contributed 0–2 points, producing a total adapted hBCS "
    "from 0 to 18. Higher scores indicated a more favorable profile. "
    "This harmonized score was adapted for the available HRS and CHARLS "
    "variables and should not be described as the official 21-point "
    "Brain Care Score."
)

supplement_path = (
    SUBMISSION /
    "Supplementary_Tables.docx"
)

supplement_document.save(
    supplement_path
)

print("Submission tables completed:")
print(f" - {main_path}")
print(f" - {supplement_path}")
